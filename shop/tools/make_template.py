"""
고객에게 보내는 상품 등록 엑셀 양식을 만듭니다.

  python3 shop/tools/make_template.py shop/tools/상품등록_양식_4종목_4000개.xlsx

- 상품목록: 4~7번째 줄 예시, 8번째 줄 안내, 9번째 줄부터 입력 (4,000줄)
- 실측사이즈: 4~7번째 줄 예시, 8번째 줄 안내, 9번째 줄부터 입력
필요한 것: pip install openpyxl
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
import sys
out=sys.argv[1]
F="맑은 고딕"
def font(**k): return Font(name=F, **k)
REQ=PatternFill("solid",fgColor="FFF2CC")   # 필수 (연노랑)
OPT=PatternFill("solid",fgColor="EAF1FB")   # 선택 (연파랑)
EX=PatternFill("solid",fgColor="EDEDED")    # 예시
HEAD=PatternFill("solid",fgColor="3A3129")
thin=Side(style="thin",color="C9C2B8"); B=Border(left=thin,right=thin,top=thin,bottom=thin)
wrap=Alignment(wrap_text=True,vertical="center")
center=Alignment(horizontal="center",vertical="center",wrap_text=True)
import os; N=int(os.environ.get("N","4000")); FIRST=9; LAST=FIRST+N-1   # 실제 입력 줄: 9~4008 (4~7 예시, 8 안내)

wb=Workbook()

# ---------------- 1. 작성 안내 ----------------
g=wb.active; g.title="작성 안내"
g.sheet_view.showGridLines=False
g.column_dimensions["A"].width=3; g.column_dimensions["B"].width=22; g.column_dimensions["C"].width=16; g.column_dimensions["D"].width=70
r=2
g.cell(r,2,"상품 등록 엑셀 양식 (의류 · 가방 · 쥬얼리 · 티셔츠 / 4,000개)").font=font(size=16,bold=True); r+=2
steps=["① '상품목록' 시트에 상품 1개당 한 줄씩 적어 주세요. (9번째 줄부터, 회색 예시 4줄은 지우지 않아도 됩니다)",
 "② 노란 칸은 꼭 채워야 하는 필수 항목, 파란 칸은 있으면 좋은 선택 항목입니다.",
 "③ 색상·사이즈가 여러 개면 쉼표(,)로 나눠 적어 주세요. 예) 블랙,화이트,네이비 / S,M,L,XL",
 "④ 사진 파일 이름은 실제 사진 파일 이름과 똑같이 적고, 사진은 폴더 하나에 모아 함께 보내 주세요.",
 "⑤ 의류·티셔츠·가방의 실측 치수는 '실측사이즈' 시트에 적어 주세요. (쥬얼리는 적지 않아도 됩니다)",
 "⑥ 맨 오른쪽 '확인' 칸에 '필수 항목 빠짐'이나 '상품코드 중복'이 뜨면 그 줄을 고쳐 주세요.",
 "⑦ '현황' 시트에서 종목별로 몇 개가 입력됐는지 볼 수 있습니다."]
for s in steps:
    g.cell(r,2,s).font=font(size=11); g.merge_cells(start_row=r,start_column=2,end_row=r,end_column=4); r+=1
r+=1
for txt,fill in [("필수 항목",REQ),("선택 항목",OPT),("예시 (지우지 않아도 됨)",EX)]:
    c=g.cell(r,2,txt); c.fill=fill; c.font=font(size=10,bold=True); c.border=B; r+=1
r+=1
g.cell(r,2,"항목 설명").font=font(size=13,bold=True); r+=1
for i,h in enumerate(["항목","필수 / 선택","적는 법"]):
    c=g.cell(r,2+i,h); c.font=font(bold=True,color="FFFFFF"); c.fill=HEAD; c.alignment=center; c.border=B
r+=1
COLS=[ # (헤더, 필수?, 너비, 설명)
 ("상품코드",True,12,"상품마다 겹치지 않는 번호. 종목 머리글자 + 번호로 적어 주세요. 의류 CL-0001, 가방 BG-0001, 쥬얼리 JW-0001, 티셔츠 TS-0001"),
 ("종목",True,10,"목록에서 고르기: 의류 / 가방 / 쥬얼리 / 티셔츠"),
 ("세부분류",False,12,"아래 '세부분류 예시' 참고. 홈페이지에서 작은 분류 버튼으로 나옵니다"),
 ("상품명",True,28,"손님에게 보이는 이름. 30자 안쪽을 권합니다"),
 ("판매가(원)",True,12,"숫자만 적어 주세요. 예) 39000"),
 ("할인가(원)",False,12,"할인할 때만 숫자로. 비워 두면 판매가로 보입니다"),
 ("색상",False,18,"쉼표로 나눠 적기. 예) 블랙,화이트"),
 ("사이즈",False,18,"쉼표로 나눠 적기. 티셔츠 S,M,L,XL / 반지 9호,11호 / 한 사이즈면 FREE"),
 ("소재",False,18,"예) 14K 골드, 925 실버, 면 100%, 소가죽"),
 ("짧은 설명",False,40,"상품 목록과 상세페이지 맨 위에 나오는 1~2줄 소개 (100자 안쪽)"),
 ("대표사진 파일명",True,22,"첫 번째로 보일 사진. 예) TS-0001_1.jpg"),
 ("추가사진 파일명",False,30,"쉼표로 나눠 적기. 예) TS-0001_2.jpg,TS-0001_3.jpg"),
 ("판매상태",False,10,"목록에서 고르기: 판매중 / 품절 / 숨김. 비워 두면 판매중"),
 ("표시",False,10,"목록에서 고르기: 신상 / 추천 / 베스트. 상품 사진에 작은 표시가 붙습니다"),
 ("메모",False,24,"저희에게 전할 말 (홈페이지에는 안 나옵니다)"),
]
for h,req,w,d in COLS:
    g.cell(r,2,h).font=font(bold=True); g.cell(r,3,"필수" if req else "선택").font=font()
    g.cell(r,3).fill=REQ if req else OPT; g.cell(r,4,d).font=font(); g.cell(r,4).alignment=wrap
    for c in range(2,5): g.cell(r,c).border=B
    g.cell(r,3).alignment=center; r+=1
r+=1
g.cell(r,2,"세부분류 예시").font=font(size=13,bold=True); r+=1
for i,h in enumerate(["종목","상품코드 머리글자","세부분류 예시"]):
    c=g.cell(r,2+i,h); c.font=font(bold=True,color="FFFFFF"); c.fill=HEAD; c.alignment=center; c.border=B
r+=1
for a,b,c_ in [("의류","CL","아우터, 셔츠, 바지, 원피스, 니트"),("가방","BG","숄더백, 토트백, 크로스백, 백팩, 클러치"),("쥬얼리","JW","반지, 목걸이, 귀걸이, 팔찌"),("티셔츠","TS","반팔, 긴팔, 민소매, 오버핏")]:
    for i,v in enumerate([a,b,c_]):
        cell=g.cell(r,2+i,v); cell.font=font(); cell.border=B; cell.alignment=wrap
    r+=1
r+=1
g.cell(r,2,"사진 보내는 법").font=font(size=13,bold=True); r+=1
for s in ["· 파일 이름은 '상품코드_번호.jpg'로 맞춰 주시면 가장 빠르게 등록됩니다. 예) JW-0001_1.jpg, JW-0001_2.jpg",
          "· 정사각형(1:1) 사진, 가로 1,000px 이상을 권합니다. JPG 또는 PNG로 보내 주세요.",
          "· 사진이 많으면 종목별 폴더(의류/가방/쥬얼리/티셔츠)로 나눠 압축해서 보내 주세요."]:
    g.cell(r,2,s).font=font(); g.merge_cells(start_row=r,start_column=2,end_row=r,end_column=4); r+=1

# ---------------- 2. 상품목록 ----------------
p=wb.create_sheet("상품목록")
ncol=len(COLS); chk=ncol+1
p.cell(1,1,"상품목록 — 상품 1개당 한 줄 (9번째 줄부터 입력)").font=font(size=14,bold=True)
p.cell(2,1,"노란 칸 = 필수 · 파란 칸 = 선택 · 회색 = 예시 · 여러 개는 쉼표(,)로 구분").font=font(size=10,color="6B6158")
for i,(h,req,w,d) in enumerate(COLS,1):
    c=p.cell(3,i,h+(" *" if req else "")); c.font=font(bold=True,color="FFFFFF"); c.fill=HEAD; c.alignment=center; c.border=B
    c.comment=Comment(d,"안내"); p.column_dimensions[c.column_letter].width=w
c=p.cell(3,chk,"확인"); c.font=font(bold=True,color="FFFFFF"); c.fill=HEAD; c.alignment=center; c.border=B
p.column_dimensions[c.column_letter].width=16
p.row_dimensions[3].height=30
examples=[
 ["예시-JW-0001","쥬얼리","반지","루미에르 다이아몬드 반지",89000,None,"화이트골드,로즈골드","9호,11호,13호","14K 골드","6발 프롱으로 감싼 라운드 스톤, 매일 끼기 좋은 높이","JW-0001_1.jpg","JW-0001_2.jpg,JW-0001_3.jpg","판매중","베스트","각인 가능"],
 ["예시-CL-0001","의류","아우터","울 블렌드 싱글 코트",159000,129000,"베이지,차콜","S,M,L","울 50% 폴리 50%","허벅지를 덮는 기장, 안감이 있어 따뜻합니다","CL-0001_1.jpg","CL-0001_2.jpg","판매중","신상",None],
 ["예시-TS-0001","티셔츠","반팔","오버핏 로고 반팔 티셔츠",35000,None,"화이트,블랙,그레이","S,M,L,XL","면 100% (20수)","탄탄한 20수 원단, 여유 있는 오버핏","TS-0001_1.jpg","TS-0001_2.jpg,TS-0001_3.jpg","판매중","추천",None],
 ["예시-BG-0001","가방","숄더백","미니 레더 숄더백",89000,None,"블랙,카멜","FREE","소가죽","스마트폰·지갑이 들어가는 크기, 끈 길이 조절","BG-0001_1.jpg",None,"품절",None,"다음 달 재입고"],
]
L=lambda i: p.cell(1,i).column_letter
def check(r):
    req=[L(i) for i,(h,rq,w,d) in enumerate(COLS,1) if rq]
    miss=",".join(f"{c}{r}=\"\"" for c in req)
    return (f'=IF(COUNTA(A{r}:{L(ncol)}{r})=0,"",IF(OR({miss}),"필수 항목 빠짐",'
            f'IF(COUNTIF($A${FIRST}:$A${LAST},A{r})>1,"상품코드 중복","OK")))')
for k,row in enumerate(examples):
    r=4+k
    for i,v in enumerate(row,1):
        c=p.cell(r,i,v); c.fill=EX; c.font=font(color="6B6158",italic=True); c.border=B; c.alignment=wrap
    c=p.cell(r,chk,"예시"); c.fill=EX; c.font=font(color="6B6158",italic=True); c.border=B; c.alignment=center
p.cell(FIRST-1,1,"▼ 여기부터 입력해 주세요").font=font(bold=True,color="A9853F")
for r in range(FIRST,LAST+1):
    for i,(h,req,w,d) in enumerate(COLS,1):
        c=p.cell(r,i); c.fill=REQ if req else OPT; c.border=B
    c=p.cell(r,chk,check(r)); c.alignment=center; c.border=B
for col in ("E","F"):
    for r in range(4,LAST+1): p[f"{col}{r}"].number_format='#,##0'
# 드롭다운
def dv(formula,col,title,msg):
    d=DataValidation(type="list",formula1=formula,allow_blank=True,showErrorMessage=True,errorTitle=title,error=msg)
    p.add_data_validation(d); d.add(f"{col}{FIRST}:{col}{LAST}")
dv('"의류,가방,쥬얼리,티셔츠"',"B","종목","의류 / 가방 / 쥬얼리 / 티셔츠 중에서 골라 주세요")
dv('"판매중,품절,숨김"',"M","판매상태","판매중 / 품절 / 숨김 중에서 골라 주세요")
dv('"신상,추천,베스트"',"N","표시","신상 / 추천 / 베스트 중에서 골라 주세요")
num=DataValidation(type="whole",operator="greaterThanOrEqual",formula1="0",allow_blank=True,showErrorMessage=True,errorTitle="가격",error="원 단위 숫자만 적어 주세요. 예) 39000")
p.add_data_validation(num); num.add(f"E{FIRST}:F{LAST}")
from openpyxl.formatting.rule import FormulaRule
p.conditional_formatting.add(f"{L(chk)}{FIRST}:{L(chk)}{LAST}",FormulaRule(formula=[f'AND({L(chk)}{FIRST}<>"",{L(chk)}{FIRST}<>"OK")'],font=Font(name=F,color="C00000",bold=True),fill=PatternFill("solid",fgColor="FDE2E2")))
p.conditional_formatting.add(f"{L(chk)}{FIRST}:{L(chk)}{LAST}",FormulaRule(formula=[f'{L(chk)}{FIRST}="OK"'],font=Font(name=F,color="2E7D32",bold=True)))
p.freeze_panes="E4"
p.auto_filter.ref=f"A3:{L(chk)}{LAST}"

# ---------------- 3. 실측사이즈 ----------------
s=wb.create_sheet("실측사이즈")
s.cell(1,1,"실측사이즈 — 사이즈 1개당 한 줄 (해당하는 칸만, 단위 cm)").font=font(size=14,bold=True)
s.cell(2,1,"의류: 어깨·가슴단면·허리·총장·소매 (바지는 허리·총장) / 티셔츠: 어깨·가슴단면·총장·소매 / 가방: 가로·세로·폭·끈길이").font=font(size=10,color="6B6158")
SH=[("상품코드",14),("사이즈",10),("어깨",10),("가슴단면",10),("허리",10),("총장",10),("소매",10),("가로",10),("세로",10),("폭",10),("끈길이",10)]
for i,(h,w) in enumerate(SH,1):
    c=s.cell(3,i,h); c.font=font(bold=True,color="FFFFFF"); c.fill=HEAD; c.alignment=center; c.border=B
    s.column_dimensions[c.column_letter].width=w
ex=[["예시-CL-0001","S",44,53,None,92,59],["예시-CL-0001","M",46,56,None,94,60],["예시-TS-0001","M",52,55,None,70,22],["예시-BG-0001","FREE",None,None,None,None,None,22,15,7,120]]
for k,row in enumerate(ex):
    for i in range(1,len(SH)+1):
        v=row[i-1] if i-1<len(row) else None
        c=s.cell(4+k,i,v); c.fill=EX; c.font=font(color="6B6158",italic=True); c.border=B; c.alignment=center
s.cell(8,1,"▼ 여기부터 입력해 주세요").font=font(bold=True,color="A9853F")
for r in range(9,1009):
    for i in range(1,len(SH)+1):
        c=s.cell(r,i); c.fill=REQ if i<=2 else OPT; c.border=B
s.freeze_panes="C4"

# ---------------- 4. 현황 ----------------
h=wb.create_sheet("현황")
h.sheet_view.showGridLines=False
h.column_dimensions["A"].width=3; h.column_dimensions["B"].width=16
for col in "CDEF": h.column_dimensions[col].width=14
h.cell(2,2,"입력 현황 (자동 계산 · 예시 줄은 세지 않음)").font=font(size=14,bold=True)
hd=["종목","입력한 상품","확인 OK","고칠 줄","품절"]
for i,t in enumerate(hd):
    c=h.cell(4,2+i,t); c.font=font(bold=True,color="FFFFFF"); c.fill=HEAD; c.alignment=center; c.border=B
rng=lambda col: f"상품목록!${col}${FIRST}:${col}${LAST}"
CK=L(chk)
for k,cat in enumerate(["의류","가방","쥬얼리","티셔츠"]):
    r=5+k
    h.cell(r,2,cat)
    h.cell(r,3,f'=COUNTIF({rng("B")},B{r})')
    h.cell(r,4,f'=COUNTIFS({rng("B")},B{r},{rng(CK)},"OK")')
    h.cell(r,5,f'=C{r}-D{r}')
    h.cell(r,6,f'=COUNTIFS({rng("B")},B{r},{rng("M")},"품절")')
h.cell(9,2,"종목 없이 입력")
h.cell(9,3,f'=COUNTA({rng("A")})-SUM(C5:C8)')
h.cell(9,4,0); h.cell(9,5,"=C9-D9"); h.cell(9,6,0)
h.cell(10,2,"합계")
for col in "CDEF": h[f"{col}10"]=f"=SUM({col}5:{col}9)"
for r in range(5,11):
    for c in range(2,7):
        cell=h.cell(r,c); cell.border=B; cell.font=font(bold=(r==10)); cell.alignment=center
        if c>2: cell.number_format='#,##0'
    if r==10:
        for c in range(2,7): h.cell(r,c).fill=PatternFill("solid",fgColor="F3EBDD")
h.cell(12,2,"계약 등록 개수").font=font(bold=True); h["C12"]=4000; h["C12"].font=font(color="0000FF"); h["C12"].number_format='#,##0'
h["D12"]="← 계약서 기준 (고객과 합의한 개수)"; h["D12"].font=font(size=9,color="6B6158")
h.cell(13,2,"남은 개수").font=font(bold=True); h["C13"]="=MAX(0,C12-C10)"; h["C13"].number_format='#,##0'
h.cell(14,2,"계약 초과").font=font(bold=True); h["C14"]="=MAX(0,C10-C12)"; h["C14"].number_format='#,##0'
h["D14"]="← 초과분은 1개당 1,000원 (계약 조건)"; h["D14"].font=font(size=9,color="6B6158")
h.cell(15,2,"진행률").font=font(bold=True); h["C15"]="=IF(C12=0,0,MIN(1,C10/C12))"; h["C15"].number_format='0.0%'
for r in range(12,16): h.cell(r,3).alignment=center; h.cell(r,3).border=B

wb.move_sheet("현황",offset=-2)  # 작성 안내, 현황, 상품목록, 실측사이즈
wb.active=0
wb.calculation.fullCalcOnLoad=True
wb.save(out)
